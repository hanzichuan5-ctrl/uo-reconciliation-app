"""Generate exports in copies. The source workbook is never opened for writing."""
from __future__ import annotations
import collections
from contextlib import closing
import copy
import json
import os
import re
import shutil
import subprocess
import tempfile
import threading
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl

from .core import APP, OUTPUT, ENTITIES, build_month, db, digest, event, key, now, receipt_template, selected_sources, workbook, cents

RUNTIME=Path('/Users/sato/.cache/codex-runtimes/codex-primary-runtime/dependencies')
NODE=RUNTIME/'node/bin/node'
EXPORT_LOCK=threading.Lock()
NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def run_author(data,output,mode,preview=False):
    work=output.parent/'.work'
    work.mkdir(parents=True,exist_ok=True)
    modules=work/'node_modules'
    if not modules.exists():
        modules.symlink_to(RUNTIME/'node/node_modules',target_is_directory=True)
    builder=work/'export.mjs'
    shutil.copyfile(APP/'monthly/export.mjs',builder)
    payload=work/f'{mode}.json'
    payload.write_text(json.dumps(dict(data,preview=preview),ensure_ascii=False))
    command=[str(NODE),str(builder),str(payload),str(output),mode]
    completed=subprocess.run(command,capture_output=True,text=True,timeout=150)
    if completed.returncode:
        raise RuntimeError('Excel生成失败：'+completed.stderr[-1500:])
    if not output.exists():
        raise RuntimeError('Excel输出未生成')


def receipt_part(z):
    root=ET.fromstring(z.read('xl/workbook.xml'))
    sheets=root.find('m:sheets',NS)
    index=next(i for i,s in enumerate(sheets) if s.attrib['name']=='回款')
    sheet=sheets[index]
    rid=sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
    target=next(r.attrib['Target'] for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels')) if r.attrib['Id']==rid)
    return (target.lstrip('/') if target.startswith('/') else 'xl/'+target),index,int(sheet.attrib['sheetId'])


def layout(source,rows):
    groups=collections.defaultdict(list)
    for r in rows:
        groups[r['entity']].append(r)
    with workbook(source) as w:
        s=w['回款']; s.reset_dimensions(); values=list(s.values)
    blocks=[]; offset=0
    for i,r in enumerate(values):
        if not r or r[0] not in ENTITIES+['东银河']:
            continue
        rn=i+1
        end=next((j+1 for j in range(i+2,len(values)) if values[j] and values[j][0]=='小计'),None)
        if not end or values[i+1][1]!='往来单位名称':
            raise ValueError('回款页分区或小计结构无法识别')
        old_capacity=end-rn-2
        if old_capacity<1:
            raise ValueError('回款分区没有可继承的格式行')
        selected=groups.pop(r[0],[])
        capacity=max(old_capacity,len(selected)); extra=capacity-old_capacity
        blocks.append({'entity':r[0],'old_title':rn,'old_start':rn+2,'old_total':end,'old_capacity':old_capacity,
                       'title':rn+offset,'start':rn+2+offset,'total':end+offset+extra,'capacity':capacity,'extra':extra,'offset':offset,'rows':selected})
        offset+=extra
    if any(groups.values()):
        raise ValueError('当前底稿缺少来款主体分区')
    return blocks


def same_receipts(source,selected):
    sig=lambda r:(r['entity'],r['payer'],r['date'],r['amountCents'])
    return collections.Counter(map(sig,receipt_template(source)))==collections.Counter(map(sig,selected))


def write_receipt_copy(source,selected,output):
    original=Path(source['path']).read_bytes()
    if digest(original)!=source['sha']:
        raise ValueError('底稿版本已变化，请重新选择')
    if same_receipts(source,selected):
        output.write_bytes(original)
        return {'changedParts':[],'otherPartsIdentical':True,'mode':'现有回款已一致，直接保留副本','count':len(selected),'amountCents':sum(r['amountCents'] for r in selected)}
    blocks=layout(source,selected)
    payload=output.parent/'.work'/'receipt-payload.xlsx'
    run_author({'blocks':blocks},payload,'receipt-payload')
    def shift(n):
        return n+sum(b['extra'] for b in blocks if n>=b['old_total'])
    def move(text,n):
        text=re.sub(r'^(<row\b[^>]*\br=")\d+',lambda m:m[1]+str(n),text)
        return re.sub(r'(<c\b[^>]*\br="[A-Z]+)\d+',lambda m:m[1]+str(n),text)
    with zipfile.ZipFile(source['path']) as z,zipfile.ZipFile(payload) as p:
        part,index,sheetid=receipt_part(z); xml=z.read(part).decode()
        original_rows={int(re.search(r'\br="(\d+)"',r)[1]):r for r in re.findall(r'<row\b[^>]*(?:/>|>.*?</row>)',xml)}
        if any(b['extra'] for b in blocks) and re.search(r'<(?:tableParts|drawing|legacyDrawing)\b',xml):
            raise ValueError('回款页含关联表格或绘图，扩容前需要适配该底稿；未修改源文件')
        rows={shift(n):move(r,shift(n)) for n,r in original_rows.items()}
        cells={c.attrib['r']:c for c in ET.fromstring(p.read(receipt_part(p)[0])).findall('.//m:sheetData/m:row/m:c',NS)}
        strings=[''.join(e.itertext()) for e in ET.fromstring(p.read('xl/sharedStrings.xml'))] if 'xl/sharedStrings.xml' in p.namelist() else []
        def cell_body(addr):
            c=cells.get(addr)
            if c is None:
                return '',None
            f=c.find('m:f',NS); v=c.find('m:v',NS)
            if f is not None:
                return '<f>'+f.text+'</f><v>'+v.text+'</v>',None
            t=c.attrib.get('t','n')
            text=strings[int(v.text)] if t=='s' else ''.join(c.find('m:is',NS).itertext()) if t=='inlineStr' else v.text if v is not None else ''
            if text=='':
                return '',None
            if addr.startswith('C'):
                return '<v>'+text+'</v>',None
            e=ET.Element('is');ET.SubElement(e,'t').text=text
            return ET.tostring(e,encoding='unicode'),'inlineStr'
        def setcell(r,col,donor):
            m=re.search(r'<c\b(?=[^>]*\br="'+('B' if col=='A' else col)+str(donor)+r'")[^>]*',original_rows[donor])
            style=re.search(r'\bs="(\d+)"',m[0]) if m else None
            body,t=cell_body(col+str(r))
            replacement=f'<c r="{col}{r}" s="{style[1] if style else 0}"'+(f' t="{t}"' if t else '')+'>'+body+'</c>'
            row=rows[r];found=re.search(r'<c\b(?=[^>]*\br="'+col+str(r)+r'")[^>]*(?:/>|>.*?</c>)',row)
            if found:
                row=row[:found.start()]+replacement+row[found.end():]
            else:
                later=next((m for m in re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"',row) if len(m[1])>1 or m[1]>col),None)
                row=row[:later.start()]+replacement+row[later.start():] if later else row.replace('</row>',replacement+'</row>')
            rows[r]=row
        for b in blocks:
            for n in range(b['old_total']+b['offset'],b['total']):
                rows[n]=move(original_rows[b['old_total']-1],n)
            for i in range(b['capacity']):
                for col in 'ABCD':
                    setcell(b['start']+i,col,b['old_start']+min(i,b['old_capacity']-1))
            setcell(b['total'],'C',b['old_total'])
        xml=re.sub(r'<sheetData>.*?</sheetData>','<sheetData>'+''.join(rows[n] for n in sorted(rows))+'</sheetData>',xml)
        xml=re.sub(r'(<dimension ref="[A-Z]+\d+:[A-Z]+)\d+("\s*/>)',lambda m:m[1]+str(max(rows))+m[2],xml)
        xml=re.sub(r'<mergeCell ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"\s*/>',lambda m:f'<mergeCell ref="{m[1]}{shift(int(m[2]))}:{m[3]}{shift(int(m[4]))}"/>',xml)
        replacements={part:xml.encode()}
        if 'xl/calcChain.xml' in z.namelist() and any(b['extra'] for b in blocks):
            # The chain can use an inherited sheet ID. Change only receipt-cell row addresses.
            current=None
            def move_chain(m):
                nonlocal current
                element=m[0]; found=re.search(r'\bi="(\d+)"',element)
                if found: current=int(found[1])
                if current==sheetid:
                    element=re.sub(r'(\br="[A-Z]+)(\d+)(")',lambda n:n[1]+str(shift(int(n[2])))+n[3],element)
                return element
            replacements['xl/calcChain.xml']=re.sub(r'<c\b[^>]*/>',move_chain,z.read('xl/calcChain.xml').decode()).encode()
        with zipfile.ZipFile(output,'w') as out:
            for info in z.infolist():
                out.writestr(copy.copy(info),replacements.get(info.filename,z.read(info.filename)))
        with zipfile.ZipFile(output) as out:
            if out.testzip() or out.namelist()!=z.namelist():
                raise ValueError('生成文件结构检查失败')
            changed=[n for n in z.namelist() if out.read(n)!=z.read(n)]
            if any(n not in replacements for n in changed):
                raise ValueError('检测到非授权工作表内容变化')
    # Read back only the receipt page; the preservation check above compares all other bytes without interpreting them.
    output_source=dict(source,path=str(output),sha=digest(output.read_bytes()))
    if not same_receipts(output_source,selected):
        raise ValueError('保存后的回款记录与纳入清单不一致')
    with workbook(output_source) as w:
        for b in blocks:
            if cents(w['回款'].cell(b['total'],3).value)!=sum(r['amountCents'] for r in b['rows']):
                raise ValueError('小计缓存验证失败')
    return {'changedParts':changed,'otherPartsIdentical':True,'mode':'仅填写回款页','count':len(selected),'amountCents':sum(r['amountCents'] for r in selected)}


def export_month(month,kind,preview=False):
    if kind not in ('income','receipt'):
        raise ValueError('未知导出类型')
    with EXPORT_LOCK:
        data=build_month(month); fingerprint=data['fingerprint']; eid=key(month,kind,fingerprint)
        with db() as con:
            previous=con.execute('SELECT * FROM exports WHERE id=?',(eid,)).fetchone()
        if previous and Path(previous['path']).exists() and digest(Path(previous['path']).read_bytes())==previous['sha']:
            return dict(previous)
        folder=OUTPUT/month/fingerprint[:10]
        folder.mkdir(parents=True,exist_ok=True)
        if kind=='income':
            if not data['income']['cases']:
                raise ValueError('本月尚无可生成清单的实际发票')
            output=folder/f'{month}_收入操作清单.xlsx'
            run_author(data,output,'income',preview)
            expected={(i['entity'],i['no']):i['amountCents'] for c in data['income']['cases'] for i in c['invoices']}
            with closing(openpyxl.load_workbook(output,read_only=True,data_only=True)) as saved:
                evidence=saved['发票依据']
                actual={(evidence.cell(r,3).value,evidence.cell(r,2).value):cents(evidence.cell(r,6).value) for r in range(5,5+len(expected))}
                if actual!=expected or len(actual)!=len(expected):
                    raise ValueError('Excel保存后的发票号码或金额校验失败')
                if cents(saved['收入操作清单'].cell(len(data['income']['cases'])+7,6).value)!=data['income']['inScopeAmountCents']:
                    raise ValueError('渠道发票合计缓存校验失败')
            verification={'invoiceCount':data['income']['invoiceCount'],'channelInvoiceCount':data['income']['inScopeCount'],
                          'channelAmountCents':data['income']['inScopeAmountCents'],'scopeReviewCount':data['income']['scopeReviewCount'],
                          'allInvoiceIdsAndAmountsVerified':True,'sourceBound':True}
        else:
            source=next((s for s in selected_sources(month) if s['role']=='template'),None)
            if not source:
                raise ValueError('本月尚未指定回款底稿')
            if data['receipts']['unmatchedTemplateCount']:
                raise ValueError('底稿中有回款未找到本月银行来源，暂不生成替换副本')
            selected=[r for r in data['receipts']['rows'] if r['status']=='included']
            if not selected:
                raise ValueError('尚无已纳入的回款，不能用空表替换底稿')
            output=folder/f'{month}_回款填写结果.xlsx'
            verification=write_receipt_copy(source,selected,output)
            verification['pendingScopeCount']=sum(r['status'] in ('review','proposed') for r in data['receipts']['rows'])
            verification['template']=source['name']; verification['templateSha']=source['sha']
            if preview:
                run_author({'path':str(output)},folder/'receipts.png','preview-template')
        if build_month(month)['fingerprint']!=fingerprint:
            raise ValueError('生成期间资料或确认发生变化，请重新生成')
        row={'id':eid,'month':month,'kind':kind,'fingerprint':fingerprint,'path':str(output),'sha':digest(output.read_bytes()),'created_at':now(),'verification':json.dumps(verification,ensure_ascii=False)}
        with db() as con:
            con.execute('INSERT OR REPLACE INTO exports VALUES(?,?,?,?,?,?,?,?)',tuple(row.values()))
            event(con,month,'生成Excel',output.name+'（副本）')
        (folder/f'{kind}-verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2))
        return row
